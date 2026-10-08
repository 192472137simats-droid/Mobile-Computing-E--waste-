"""Build the VESTIGE evaluation deck from genuine local web screenshots.

Run: python docs/submission/make_presentation.py
Requires: python-pptx. The browser captures in screenshots/ are source assets.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
SHOTS = HERE / "screenshots"
HERO = ROOT / "frontend" / "src" / "assets" / "vestige_hero.jpg"
OUT = HERE / "VESTIGE_Evaluation_Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

CREAM = RGBColor(247, 239, 221)
PAPER = RGBColor(255, 251, 242)
INK = RGBColor(31, 22, 17)
BROWN = RGBColor(48, 29, 19)
FOREST = RGBColor(22, 75, 51)
GOLD = RGBColor(181, 132, 54)
LIGHT_GOLD = RGBColor(231, 206, 151)
MUTED = RGBColor(112, 96, 74)
WHITE = RGBColor(255, 255, 255)


def rect(slide, x, y, w, h, color, rounded=False, border=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if border:
        shape.line.color.rgb = border
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, content, x, y, w, h, size=20, color=INK, bold=False,
         serif=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = content
    p.alignment = align
    p.font.name = "Georgia" if serif else "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def image_cover(slide, path, x, y, w, h):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return pic


def shot(slide, filename, x, y, w, h, caption=None):
    rect(slide, x - 0.08, y - 0.08, w + 0.16, h + 0.16, WHITE, rounded=True, border=LIGHT_GOLD)
    image_cover(slide, SHOTS / filename, x, y, w, h)
    if caption:
        rect(slide, x, y + h - 0.45, w, 0.45, BROWN)
        text(slide, caption, x + 0.17, y + h - 0.41, w - 0.34, 0.35, 13, WHITE, True)


def slide(title, chapter, number, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BROWN if dark else CREAM
    rect(s, 0, 0, 13.333, 0.075, GOLD)
    text(s, chapter.upper(), 0.62, 0.34, 7.5, 0.28, 11, LIGHT_GOLD if dark else GOLD, True)
    text(s, title, 0.62, 0.75, 12.0, 0.72, 29, WHITE if dark else INK, True, serif=True)
    rect(s, 0.62, 7.02, 12.04, 0.012, GOLD if dark else LIGHT_GOLD)
    text(s, "VESTIGE  /  EVALUATION PRESENTATION", 0.63, 7.08, 9.0, 0.19, 9, LIGHT_GOLD if dark else MUTED, True)
    text(s, f"{number:02}", 12.16, 7.07, 0.45, 0.22, 10, LIGHT_GOLD if dark else MUTED, True, align=PP_ALIGN.RIGHT)
    return s


def card(s, x, y, w, h, kicker, heading, body, fill=PAPER):
    rect(s, x, y, w, h, fill, rounded=True, border=LIGHT_GOLD)
    text(s, kicker.upper(), x + 0.23, y + 0.23, w - 0.46, 0.25, 10, GOLD, True)
    text(s, heading, x + 0.23, y + 0.59, w - 0.46, 0.55, 20, INK, True, serif=True)
    text(s, body, x + 0.23, y + 1.26, w - 0.46, h - 1.46, 15, MUTED, valign=MSO_ANCHOR.TOP)


def add_fade(s):
    transition = OxmlElement("p:transition")
    transition.set("spd", "med")
    transition.set("advClick", "1")
    transition.append(OxmlElement("p:fade"))
    s._element.append(transition)


# 01 — cover
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = BROWN
image_cover(s, HERO, 5.85, 0, 7.483, 7.5)
rect(s, 0, 0, 6.12, 7.5, BROWN)
rect(s, 0.72, 0.78, 0.82, 0.06, GOLD)
text(s, "VESTIGE", 0.71, 1.29, 5.0, 1.08, 50, WHITE, True, serif=True)
text(s, "GIVE TECHNOLOGY\nA SECOND LIFE", 0.73, 2.54, 4.95, 1.45, 28, LIGHT_GOLD, True)
text(s, "One connected journey for assessment, repair, recycling and rewards.", 0.73, 4.45, 4.72, 1.0, 22, WHITE)
rect(s, 0.74, 6.33, 4.85, 0.04, GOLD)
text(s, "EVALUATION DECK  •  2026", 0.74, 6.55, 4.8, 0.34, 12, LIGHT_GOLD, True)

# 02 — problem
s = slide("A broken journey for old electronics", "01 / The problem", 2)
image_cover(s, HERO, 7.45, 1.66, 5.18, 4.91)
rect(s, 0.70, 1.69, 6.35, 4.85, PAPER, rounded=True)
text(s, "THE DECISION GAP", 1.00, 2.00, 5.5, 0.35, 13, GOLD, True)
text(s, "Repair, recycle, or leave it unused?", 1.00, 2.51, 5.60, 1.18, 29, INK, True, serif=True)
text(s, "Owners face uncertain repair value, scattered providers and weak follow up. Vendors receive fragmented requests. Administrators lack a joined view of outcomes.", 1.00, 4.07, 5.50, 1.55, 19, MUTED)
text(s, "VESTIGE connects those steps in one traceable workflow.", 1.00, 5.89, 5.55, 0.38, 16, FOREST, True)

# 03 — platform image
s = slide("A product built around the decision", "02 / The experience", 3)
shot(s, "landing.png", 0.77, 1.77, 11.79, 4.85, "LIVE CAPTURE  •  VESTIGE LANDING PAGE")

# 04 — three roles
s = slide("Three perspectives. One shared record.", "03 / Stakeholders", 4)
card(s, 0.71, 1.71, 3.91, 4.91, "Owner", "Understand the next step", "Assess the device, book a service, track its progress and view eligible rewards.")
card(s, 4.72, 1.71, 3.91, 4.91, "Vendor", "Manage assigned work", "See structured intake details and update repair or recycling request status.")
card(s, 8.73, 1.71, 3.91, 4.91, "Administrator", "Oversee the system", "Review users, vendors, bookings, payments, rewards and platform activity.")

# 05 — user dashboard
s = slide("Owner portal: the circular journey starts here", "04 / Owner experience", 5)
shot(s, "user_dashboard.png", 0.75, 1.75, 8.55, 4.87, "LIVE CAPTURE  •  USER DASHBOARD")
rect(s, 9.58, 1.75, 3.02, 4.87, BROWN, rounded=True)
text(s, "01", 9.89, 2.12, 1.2, 0.55, 29, LIGHT_GOLD, True, serif=True)
text(s, "One entry point", 9.89, 2.90, 2.40, 0.85, 24, WHITE, True, serif=True)
text(s, "The dashboard connects diagnostics, collection, bookings and rewards.", 9.89, 4.00, 2.38, 1.62, 17, WHITE)

# 06 — assessment
s = slide("Transparent, rules based assessment", "05 / Decision support", 6)
shot(s, "assessment.png", 4.79, 1.71, 7.82, 4.92, "LIVE CAPTURE  •  DEVICE ASSESSMENT")
rect(s, 0.71, 1.71, 3.78, 4.92, FOREST, rounded=True)
text(s, "INPUTS", 1.00, 2.09, 3.2, 0.36, 13, LIGHT_GOLD, True)
text(s, "Age\nCondition\nRepair cost / value\nParts availability\nKnown faults", 1.00, 2.66, 3.12, 2.55, 20, WHITE)
rect(s, 1.00, 5.58, 2.90, 0.02, LIGHT_GOLD)
text(s, "Score 5–95  ·  Repair at ≥50", 1.00, 5.78, 3.08, 0.48, 15, WHITE, True)

# 07 — workflow
s = slide("From assessment to a recorded outcome", "06 / Workflow", 7)
flow = [
    ("01", "ASSESS", "Enter details and receive a reasoned suggestion."),
    ("02", "SUBMIT", "Save the device and request repair or recycling."),
    ("03", "PROCESS", "Vendor manages the assigned request and status."),
    ("04", "REWARD", "Owner sees applicable points and history."),
]
for i, (n, title, body) in enumerate(flow):
    x = 0.71 + i * 3.18
    rect(s, x, 2.05, 2.92, 3.63, BROWN if i == 3 else PAPER, rounded=True, border=GOLD)
    text(s, n, x + 0.22, 2.30, 1.1, 0.55, 30, LIGHT_GOLD if i == 3 else GOLD, True, serif=True)
    text(s, title, x + 0.22, 3.20, 2.48, 0.49, 18, WHITE if i == 3 else INK, True)
    text(s, body, x + 0.22, 3.91, 2.40, 1.20, 16, WHITE if i == 3 else MUTED)
text(s, "Recommendation = decision support. A technician must confirm the actual service route.", 0.82, 6.15, 11.8, 0.42, 16, FOREST, True)

# 08 — vendor
s = slide("Vendor portal: work stays visible", "07 / Service operations", 8)
shot(s, "vendor_dashboard.png", 0.75, 1.75, 8.55, 4.87, "LIVE CAPTURE  •  VENDOR DASHBOARD")
rect(s, 9.58, 1.75, 3.02, 4.87, FOREST, rounded=True)
text(s, "02", 9.89, 2.12, 1.2, 0.55, 29, LIGHT_GOLD, True, serif=True)
text(s, "Assigned work", 9.89, 2.90, 2.40, 0.85, 24, WHITE, True, serif=True)
text(s, "Repair and recycling requests appear with customer details, status and next actions.", 9.89, 4.00, 2.38, 1.68, 17, WHITE)

# 09 — admin
s = slide("Administrator portal: operational oversight", "08 / Governance", 9)
shot(s, "admin_dashboard.png", 0.75, 1.75, 8.55, 4.87, "LIVE CAPTURE  •  ADMIN DASHBOARD")
rect(s, 9.58, 1.75, 3.02, 4.87, BROWN, rounded=True)
text(s, "03", 9.89, 2.12, 1.2, 0.55, 29, LIGHT_GOLD, True, serif=True)
text(s, "Joined view", 9.89, 2.90, 2.40, 0.85, 24, WHITE, True, serif=True)
text(s, "User, vendor, request, reward and payment records support review and follow up.", 9.89, 4.00, 2.38, 1.68, 17, WHITE)

# 10 — architecture
s = slide("The system behind the experience", "09 / Architecture", 10)
layers = [
    (0.79, "WEB EXPERIENCE", "React 19  •  Vite 8  •  role based routes", BROWN),
    (4.91, "APPLICATION API", "Java 21  •  Spring Boot 3.4.3  •  JWT", FOREST),
    (9.03, "DATA + SERVICES", "H2 demo  •  PostgreSQL schema  •  Razorpay", BROWN),
]
for x, head, body, color in layers:
    rect(s, x, 1.98, 3.55, 2.15, color, rounded=True)
    text(s, head, x + 0.24, 2.30, 3.05, 0.38, 14, LIGHT_GOLD, True)
    text(s, body, x + 0.24, 2.85, 3.04, 0.90, 18, WHITE)
text(s, "→", 4.44, 2.69, 0.41, 0.6, 31, GOLD, True)
text(s, "→", 8.56, 2.69, 0.41, 0.6, 31, GOLD, True)
rect(s, 0.79, 4.55, 11.79, 1.55, PAPER, rounded=True, border=LIGHT_GOLD)
text(s, "SECURITY + TRACEABILITY", 1.13, 4.83, 3.52, 0.34, 15, GOLD, True)
text(s, "Signed tokens, role and ownership checks, persisted request status and administrative audit actions.", 1.13, 5.28, 11.04, 0.5, 19, INK)

# 11 — live evidence collage
s = slide("What is verified in the local build", "10 / Evidence", 11)
shot(s, "landing.png", 0.78, 1.77, 5.60, 3.63, "LIVE LANDING PAGE")
shot(s, "assessment.png", 6.70, 1.77, 5.86, 3.63, "LIVE ASSESSMENT FORM")
rect(s, 0.78, 5.73, 11.78, 0.77, FOREST, rounded=True)
text(s, "57 backend tests passed   •   frontend build passed   •   lint exited successfully", 1.08, 5.87, 11.20, 0.43, 20, WHITE, True, align=PP_ALIGN.CENTER)

# 12 — boundaries
s = slide("Honest boundaries of the prototype", "11 / Readiness", 12)
items = [
    ("LIVE PAYMENTS", "Gateway credentials and a real checkout verification are pending."),
    ("PRODUCTION DATA", "PostgreSQL migrations and deployed services need target environment checks."),
    ("IMPACT ESTIMATES", "Device mass and CO₂e values are assumptions, not measured results."),
    ("FIELD OPERATIONS", "Vendor qualification, pickup and recycling proof need real processes."),
]
for i, (head, body) in enumerate(items):
    y = 1.78 + i * 1.17
    rect(s, 0.79, y, 11.78, 0.96, PAPER if i % 2 == 0 else RGBColor(241, 227, 196), rounded=True)
    text(s, head, 1.06, y + 0.17, 3.05, 0.57, 16, FOREST, True)
    text(s, body, 4.03, y + 0.17, 8.14, 0.57, 17, INK)

# 13 — evaluation path
s = slide("The next step: validate with real outcomes", "12 / Roadmap", 13)
card(s, 0.72, 1.82, 3.88, 4.63, "01 / Pilot", "Observe real jobs", "Test intake and status tracking with repair and recycling partners.")
card(s, 4.73, 1.82, 3.88, 4.63, "02 / Calibrate", "Improve the rules", "Tune assessment thresholds against actual diagnoses and repair costs.")
card(s, 8.74, 1.82, 3.88, 4.63, "03 / Verify", "Document impact", "Add chain of custody and traceable, device specific environmental methods.")

# 14 — closing
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = BROWN
image_cover(s, HERO, 6.54, 0, 6.793, 7.5)
rect(s, 0, 0, 6.71, 7.5, BROWN)
text(s, "VESTIGE", 0.76, 1.08, 5.34, 0.92, 45, WHITE, True, serif=True)
text(s, "ASSESS  →  REPAIR  →\nREUSE  →  RECYCLE  →\nREWARD", 0.78, 2.35, 5.2, 2.00, 24, LIGHT_GOLD, True)
rect(s, 0.78, 5.03, 4.91, 0.035, GOLD)
text(s, "Problem statement + full project details in the repository", 0.78, 5.35, 5.09, 0.70, 17, WHITE)
text(s, "github.com/Jaswanth300/vestige", 0.78, 6.36, 5.23, 0.42, 16, LIGHT_GOLD, True)

for page in prs.slides:
    add_fade(page)

prs.save(OUT)
print(f"Created {OUT} with {len(prs.slides)} slides and fade transitions")
